"""템플릿 기반 엑셀 생성기 (FR-M3-1~8, 스펙 5.3/9.3절 v1 초안 레이아웃).

조직 기존 보고/수익자료 템플릿 원본을 아직 확보하지 못해(스펙 11절
오픈이슈), 확보 전까지는 9.3절의 v1 초안 레이아웃으로 우선 구현한다.
템플릿 구조 변경에 대비해 시트/컬럼 정의는 이 파일 상단 상수로 분리해
두었다 — 나중에 실제 템플릿이 오면 이 부분만 바꾸면 된다.

레이아웃:
  - 시트: "신규법인"(맨 앞) + "1팀"~"6팀"
  - 컬럼: 법인명 / 키워드 / 영업시그널 카테고리 / 이슈요약 / 원문링크
  - 본문 글자 크기 10pt, 이슈요약 칸은 줄바꿈(wrap_text) 적용
  - 법인명/키워드/영업시그널 카테고리/원문링크는 왼쪽 가운데 정렬
  - 기준일자 헤더 1회 표기
  - 법인명 그룹핑. 그룹은 보도량(대표 기사 + 같은 사건으로 묶인 관련 기사
    수)이 많은 순, 같으면 최신 기사 발행일 순으로 배치한다. 기사 25건이
    나온 증자와 1건짜리 단신이 같은 무게로 섞이지 않게, 많이 보도된 사건이
    위로 오게 한다. 그룹 내부도 보도량 많은 순, 같으면 발행일 최신순

파일명: "이슈트래커(YYMMDD).xlsx". 연휴처럼 여러 날짜를 하나로 묶어 낼
때는 issues 파일의 run_date를 "2026-09-19~2026-09-20"처럼 적어두면
"이슈트래커(260919~260920).xlsx"로 나온다.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "module2_news_pipeline"))
from date_utils import KST, month_folder_name, parse_pub_date  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ISSUES_DIR = REPO_ROOT / "module2_news_pipeline" / "output"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "outputs"

TEAM_SHEET_NAMES = [f"{i}팀" for i in range(1, 7)]
NEW_CORP_SHEET_NAME = "신규법인"
COLUMNS = ["법인명", "키워드", "영업시그널 카테고리", "이슈요약", "원문링크"]
COLUMN_WIDTHS = [22, 22, 24, 65, 45]
BODY_FONT_SIZE = 10
SUMMARY_COLUMN_INDEX = COLUMNS.index("이슈요약") + 1  # 1-based
LEFT_CENTER_COLUMNS = ["법인명", "키워드", "영업시그널 카테고리", "원문링크"]
LEFT_CENTER_COLUMN_INDEXES = [COLUMNS.index(name) + 1 for name in LEFT_CENTER_COLUMNS]
_EPOCH = datetime(1970, 1, 1, tzinfo=KST)


def _pub_date(record: dict) -> datetime:
    return parse_pub_date(record.get("pub_date", "")) or _EPOCH


def _coverage(record: dict) -> int:
    """대표 기사 1건 + 같은 사건으로 묶인 관련 기사 수."""
    return 1 + len(record.get("related_articles") or [])


def _record_order_key(record: dict) -> tuple[int, datetime]:
    return (_coverage(record), _pub_date(record))


def _group_by_corp(records: list[dict]) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = {}
    for record in records:
        grouped.setdefault(record["matched_corp"], []).append(record)
    return grouped


def _ordered_corp_names(grouped: dict[str, list[dict]]) -> list[str]:
    return sorted(
        grouped.keys(),
        key=lambda corp: max(_record_order_key(r) for r in grouped[corp]),
        reverse=True,
    )


def _write_sheet(ws: Worksheet, run_date: str, grouped: dict[str, list[dict]]) -> None:
    ws.append([f"기준일자: {run_date}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    ws["A1"].font = Font(bold=True, size=BODY_FONT_SIZE)

    ws.append(COLUMNS)
    for cell in ws[2]:
        cell.font = Font(bold=True, size=BODY_FONT_SIZE)
        if cell.column in LEFT_CENTER_COLUMN_INDEXES:
            cell.alignment = Alignment(horizontal="left", vertical="center")

    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = width

    # 헤더(2행)를 고정해서 스크롤해도 항상 보이게 한다.
    ws.freeze_panes = "A3"

    if not grouped:
        ws.append(["해당 없음", "", "", "", ""])
        for cell in ws[ws.max_row]:
            cell.font = Font(size=BODY_FONT_SIZE)
    else:
        for corp_name in _ordered_corp_names(grouped):
            records = sorted(grouped[corp_name], key=_record_order_key, reverse=True)
            for record in records:
                link = record.get("originallink") or record.get("link", "")
                keywords = ", ".join(record.get("matched_keywords") or [])
                row = [corp_name, keywords, record.get("keyword_tag", ""), record.get("summary", ""), link]
                ws.append(row)
                row_idx = ws.max_row

                for cell in ws[row_idx]:
                    cell.font = Font(size=BODY_FONT_SIZE)
                    if cell.column in LEFT_CENTER_COLUMN_INDEXES:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

                # 이슈요약은 길어서 줄바꿈해야 한 눈에 읽힌다.
                summary_cell = ws.cell(row=row_idx, column=SUMMARY_COLUMN_INDEX)
                summary_cell.alignment = Alignment(wrap_text=True, vertical="top")

                link_col = len(COLUMNS)
                if link:
                    link_cell = ws.cell(row=row_idx, column=link_col)
                    link_cell.hyperlink = link
                    link_cell.font = Font(color="0563C1", underline="single", size=BODY_FONT_SIZE)

    # 담당자가 법인명/카테고리 등으로 바로 필터·정렬할 수 있도록 자동 필터를 켠다.
    last_col_letter = ws.cell(row=2, column=len(COLUMNS)).column_letter
    ws.auto_filter.ref = f"A2:{last_col_letter}{ws.max_row}"


def _short_date(date_str: str) -> str:
    """"2026-09-19" -> "260919". 파일명을 짧게 유지하기 위한 표기."""
    return datetime.strptime(date_str.strip(), "%Y-%m-%d").strftime("%y%m%d")


def _output_stem(run_date: str) -> str:
    """연휴처럼 여러 날짜를 하나로 묶은 경우 "260919~260920"처럼 범위로 표기한다."""
    if "~" in run_date:
        start, end = run_date.split("~", 1)
        return f"{_short_date(start)}~{_short_date(end)}"
    return _short_date(run_date)


def generate(issues_path: Path, output_dir: Path) -> Path:
    data = json.loads(issues_path.read_text(encoding="utf-8"))
    run_date = data.get("run_date", "")
    issues_by_corp: dict[str, list[dict]] = data.get("issues_by_corp", {})

    new_corp_records: list[dict] = []
    team_records: dict[str, list[dict]] = {name: [] for name in TEAM_SHEET_NAMES}

    for records in issues_by_corp.values():
        for record in records:
            if record.get("corp_status") == "신규":
                new_corp_records.append(record)
                continue
            for team in record.get("team") or []:
                sheet_name = f"{team}팀"
                if sheet_name in team_records:
                    team_records[sheet_name].append(record)

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet(NEW_CORP_SHEET_NAME)
    _write_sheet(ws, run_date, _group_by_corp(new_corp_records))

    for sheet_name in TEAM_SHEET_NAMES:
        ws = wb.create_sheet(sheet_name)
        _write_sheet(ws, run_date, _group_by_corp(team_records[sheet_name]))

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"이슈트래커({_output_stem(run_date)}).xlsx"
    wb.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--date",
        required=True,
        help="issues_<날짜>.json의 날짜. 합본은 2026-09-19~2026-09-21처럼 기간으로 지정",
    )
    parser.add_argument("--issues-dir", type=Path, default=DEFAULT_ISSUES_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()

    # pipeline.py/merge_llm_results.py가 "2026년 9월"처럼 월 단위 폴더에
    # 쓰므로 같은 폴더에서 찾는다.
    # 합본(기간)은 첫 날짜가 속한 달 폴더에 있다.
    month = month_folder_name(args.date.split("~", 1)[0])
    issues_path = args.issues_dir / month / f"issues_{args.date}.json"
    output_dir = args.output_dir / month
    result = generate(issues_path, output_dir)
    print(f"엑셀: {result}")


if __name__ == "__main__":
    main()
