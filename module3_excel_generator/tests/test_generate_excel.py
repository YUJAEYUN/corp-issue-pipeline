import json

from openpyxl import load_workbook

from generate_excel import NEW_CORP_SHEET_NAME, TEAM_SHEET_NAMES, generate


def _write_issues(tmp_path, records_by_corp):
    path = tmp_path / "issues_2026-09-16.json"
    path.write_text(
        json.dumps(
            {
                "run_date": "2026-09-16",
                "collected_date_from": "2026-09-15",
                "collected_date_to": "2026-09-15",
                "issues_by_corp": records_by_corp,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return path


def test_output_filename_uses_short_date_range_for_multi_day_run(tmp_path):
    path = tmp_path / "issues_weekend.json"
    path.write_text(
        json.dumps(
            {
                "run_date": "2026-09-19~2026-09-20",
                "collected_date_from": "2026-09-19",
                "collected_date_to": "2026-09-20",
                "issues_by_corp": {},
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    output_path = generate(path, tmp_path / "out")

    assert output_path.name == "이슈트래커(260919~260920).xlsx"


def test_generates_all_expected_sheets_in_order(tmp_path):
    issues_path = _write_issues(tmp_path, {})
    output_path = generate(issues_path, tmp_path / "out")

    assert output_path.name == "이슈트래커(260916).xlsx"
    wb = load_workbook(output_path)
    assert wb.sheetnames == [NEW_CORP_SHEET_NAME] + TEAM_SHEET_NAMES


def test_known_corp_issue_lands_in_its_team_sheet_with_keywords_and_link(tmp_path):
    records_by_corp = {
        "가나다전자(주)": [
            {
                "article_id": "1",
                "title": "제목",
                "summary": "가나다전자가 유상증자를 결정했다",
                "originallink": "http://example.com/1",
                "link": "http://news.naver.com/1",
                "pub_date": "Wed, 16 Sep 2026 09:00:00 +0900",
                "matched_keywords": ["유상증자", "M&A"],
                "matched_corp": "가나다전자(주)",
                "corp_status": "기존",
                "team": ["1", "3"],
                "keyword_tag": "조달 니즈 발생",
            }
        ]
    }
    issues_path = _write_issues(tmp_path, records_by_corp)
    output_path = generate(issues_path, tmp_path / "out")

    wb = load_workbook(output_path)
    for sheet_name in ["1팀", "3팀"]:
        ws = wb[sheet_name]
        row = [cell.value for cell in ws[3]]
        assert row[0] == "가나다전자(주)"
        assert row[1] == "유상증자, M&A"
        assert row[4] == "http://example.com/1"
        assert ws.cell(row=3, column=5).hyperlink is not None

    # 소속 안 된 팀 시트에는 안 들어가야 한다
    ws_untouched = wb["2팀"]
    assert [cell.value for cell in ws_untouched[3]][:1] == ["해당 없음"]


def test_new_corp_issue_lands_in_new_corp_sheet_not_team_sheets(tmp_path):
    records_by_corp = {
        "미상법인": [
            {
                "article_id": "2",
                "title": "제목",
                "summary": "신규 법인 관련 기사",
                "originallink": "http://example.com/2",
                "link": "http://news.naver.com/2",
                "pub_date": "Wed, 16 Sep 2026 09:00:00 +0900",
                "matched_keywords": ["M&A"],
                "matched_corp": "미상법인",
                "corp_status": "신규",
                "team": [],
                "keyword_tag": "기업 이벤트",
            }
        ]
    }
    issues_path = _write_issues(tmp_path, records_by_corp)
    output_path = generate(issues_path, tmp_path / "out")

    wb = load_workbook(output_path)
    ws = wb[NEW_CORP_SHEET_NAME]
    assert [cell.value for cell in ws[3]][0] == "미상법인"
    for sheet_name in TEAM_SHEET_NAMES:
        assert [cell.value for cell in wb[sheet_name][3]][:1] == ["해당 없음"]


def test_header_row_has_autofilter_and_date(tmp_path):
    issues_path = _write_issues(tmp_path, {})
    output_path = generate(issues_path, tmp_path / "out")

    wb = load_workbook(output_path)
    ws = wb[NEW_CORP_SHEET_NAME]
    assert ws["A1"].value == "기준일자: 2026-09-16"
    assert [cell.value for cell in ws[2]] == [
        "법인명",
        "키워드",
        "영업시그널 카테고리",
        "이슈요약",
        "원문링크",
    ]
    assert ws.auto_filter.ref is not None


def test_summary_column_wraps_text_and_body_uses_10pt_font(tmp_path):
    records_by_corp = {
        "가나다전자(주)": [
            {
                "article_id": "1",
                "title": "제목",
                "summary": "가나다전자가 유상증자를 결정했다. 조달자금은 시설투자에 쓸 예정이다.",
                "originallink": "http://example.com/1",
                "link": "http://news.naver.com/1",
                "pub_date": "Wed, 16 Sep 2026 09:00:00 +0900",
                "matched_keywords": ["유상증자"],
                "matched_corp": "가나다전자(주)",
                "corp_status": "기존",
                "team": ["1"],
                "keyword_tag": "조달 니즈 발생",
            }
        ]
    }
    issues_path = _write_issues(tmp_path, records_by_corp)
    output_path = generate(issues_path, tmp_path / "out")

    wb = load_workbook(output_path)
    ws = wb["1팀"]

    summary_cell = ws.cell(row=3, column=4)  # 법인명/키워드/카테고리/이슈요약
    assert summary_cell.value.startswith("가나다전자가 유상증자를")
    assert summary_cell.alignment.wrap_text is True

    for cell in ws[3]:
        assert cell.font.size == 10
    for cell in ws[2]:
        assert cell.font.size == 10


def test_left_center_alignment_on_non_summary_columns(tmp_path):
    records_by_corp = {
        "가나다전자(주)": [
            {
                "article_id": "1",
                "title": "제목",
                "summary": "가나다전자가 유상증자를 결정했다.",
                "originallink": "http://example.com/1",
                "link": "http://news.naver.com/1",
                "pub_date": "Wed, 16 Sep 2026 09:00:00 +0900",
                "matched_keywords": ["유상증자"],
                "matched_corp": "가나다전자(주)",
                "corp_status": "기존",
                "team": ["1"],
                "keyword_tag": "조달 니즈 발생",
            }
        ]
    }
    issues_path = _write_issues(tmp_path, records_by_corp)
    output_path = generate(issues_path, tmp_path / "out")

    wb = load_workbook(output_path)
    ws = wb["1팀"]

    # 법인명(1) / 키워드(2) / 카테고리(3) / 원문링크(5)는 왼쪽 가운데 정렬
    for col in (1, 2, 3, 5):
        for row in (2, 3):
            cell = ws.cell(row=row, column=col)
            assert cell.alignment.horizontal == "left"
            assert cell.alignment.vertical == "center"

    # 이슈요약(4)은 줄바꿈용 정렬(왼쪽 가운데가 아님)이라 제외
    summary_cell = ws.cell(row=3, column=4)
    assert summary_cell.alignment.vertical == "top"


def test_main_accepts_date_range_and_reads_first_month_folder(tmp_path, monkeypatch):
    import sys

    import generate_excel

    month_dir = tmp_path / "issues" / "2026년 9월"
    month_dir.mkdir(parents=True)
    (month_dir / "issues_2026-09-30~2026-10-02.json").write_text(
        json.dumps({"run_date": "2026-09-30~2026-10-02", "issues_by_corp": {}}, ensure_ascii=False),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "generate_excel.py",
            "--date",
            "2026-09-30~2026-10-02",
            "--issues-dir",
            str(tmp_path / "issues"),
            "--output-dir",
            str(tmp_path / "out"),
        ],
    )

    generate_excel.main()

    assert (tmp_path / "out" / "2026년 9월" / "이슈트래커(260930~261002).xlsx").exists()


def test_rows_are_ordered_by_coverage_then_date(tmp_path):
    def record(corp, article_id, pub_date, related=0):
        return {
            "article_id": article_id,
            "title": "제목",
            "summary": f"{corp} 요약 {article_id}",
            "originallink": f"http://example.com/{article_id}",
            "link": "",
            "pub_date": pub_date,
            "matched_keywords": ["유상증자"],
            "matched_corp": corp,
            "corp_status": "기존",
            "team": ["1"],
            "keyword_tag": "조달 니즈 발생",
            "related_articles": [{"article_id": f"{article_id}-{i}"} for i in range(related)],
        }

    records_by_corp = {
        "최신단신": [record("최신단신", "a", "Wed, 16 Sep 2026 18:00:00 +0900")],
        "대형사건": [
            record("대형사건", "b", "Wed, 16 Sep 2026 08:00:00 +0900", related=24),
            record("대형사건", "c", "Wed, 16 Sep 2026 17:00:00 +0900"),
        ],
    }
    output_path = generate(_write_issues(tmp_path, records_by_corp), tmp_path / "out")

    ws = load_workbook(output_path)["1팀"]
    rows = [(ws.cell(row=r, column=1).value, ws.cell(row=r, column=4).value) for r in range(3, 6)]
    assert rows == [
        ("대형사건", "대형사건 요약 b"),
        ("대형사건", "대형사건 요약 c"),
        ("최신단신", "최신단신 요약 a"),
    ]
